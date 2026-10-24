"""Generator HISTORY_EXCEL + task run_export (exports.md §3, §4)."""
import tempfile
from datetime import datetime, timedelta, timezone as dt_timezone
from decimal import Decimal
from unittest import mock

from django.contrib.auth import get_user_model
from django.test import TestCase, override_settings
from openpyxl import load_workbook

from apps.alerting.models import Alert
from apps.inventory.models import SparePart, StockMovement
from apps.ml_inference.models import Prediction
from apps.registry.models import Bearing, Machine
from apps.scheduling.models import WorkOrder

from .excel_report import build_history_excel
from .models import ExportJob
from .storage import output_path
from .tasks import run_export

User = get_user_model()
T0 = datetime(2026, 9, 1, tzinfo=dt_timezone.utc)
T_END = T0 + timedelta(days=10)
FORMULA = '=HYPERLINK("http://contoh.invalid","klik")'


def iso(dt):
    return dt.strftime("%Y-%m-%dT%H:%M:%SZ")


def attnames(model):
    return [f.attname for f in model._meta.concrete_fields]


def trim(row):
    """openpyxl melebarkan setiap baris ke lebar sheet dengan None; buang None di ekor."""
    row = list(row)
    while row and row[-1] is None:
        row.pop()
    return row


class ExcelBase(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.user = User.objects.create_user(username="xl_eng", password="x", role="ENGINEER")
        cls.m1 = Machine.objects.create(name="Xl Rig 1", location="L", machine_type="t", criticality="LOW")
        cls.m2 = Machine.objects.create(name="Xl Rig 2", location="L", machine_type="t", criticality="LOW")
        cls.b1 = Bearing.objects.create(machine=cls.m1, position_label="de", bearing_model="X", status="ACTIVE")
        cls.b2 = Bearing.objects.create(machine=cls.m2, position_label="de", bearing_model="X", status="ACTIVE")

        def pred(bearing, when):
            return Prediction.objects.create(
                bearing=bearing, prediction_type="CLASSIFICATION", fault_class="OUTER_RACE",
                classification_confidence=Decimal("0.950"), model_version="xl-v1", predicted_at=when)

        cls.p_start = pred(cls.b1, T0)                               # tepat di start -> ikut
        cls.p_in1 = pred(cls.b1, T0 + timedelta(days=1))
        cls.p_in2 = pred(cls.b2, T0 + timedelta(days=2))
        cls.p_before = pred(cls.b1, T0 - timedelta(days=1))
        cls.p_at_end = pred(cls.b1, T_END)                           # tepat di end -> TIDAK ikut
        cls.p_after = pred(cls.b1, T0 + timedelta(days=11))

        def alert(bearing, atype, when, **kw):
            return Alert.objects.create(bearing=bearing, alert_type=atype, severity="HIGH",
                                        triggered_at=when, last_seen_at=when, **kw)

        cls.a_in = alert(cls.b1, "FAULT_DETECTED", T0 + timedelta(days=1))
        cls.a_null = alert(None, "LOW_STOCK", T0 + timedelta(days=1))
        cls.a_out = alert(cls.b1, "FAULT_DETECTED", T0 - timedelta(days=5))

        cls.wo_in = WorkOrder.objects.create(bearing=cls.b1, status="COMPLETED", priority="HIGH", notes=FORMULA)
        cls.wo_out = WorkOrder.objects.create(bearing=cls.b1, status="DRAFT", priority="LOW")
        WorkOrder.objects.filter(pk=cls.wo_in.pk).update(created_at=T0 + timedelta(days=3))
        WorkOrder.objects.filter(pk=cls.wo_out.pk).update(created_at=T0 - timedelta(days=30))

        cls.part = SparePart.objects.create(sku="XL-1", name="Bearing XL", quantity_on_hand=3, reorder_point=2, reorder_quantity=5)
        # stock_movements.timestamp NOT NULL tanpa nilai otomatis -> wajib diisi eksplisit
        cls.mv_in = StockMovement.objects.create(
            spare_part=cls.part, movement_type="RECEIPT", quantity=5, timestamp=T0 + timedelta(days=1))
        cls.mv_out = StockMovement.objects.create(
            spare_part=cls.part, movement_type="RECEIPT", quantity=1, timestamp=T0 - timedelta(days=9))

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.enterContext(override_settings(RAW_WAVEFORM_STORAGE_PATH=self.tmp.name))

    def job(self, **extra):
        return ExportJob.objects.create(
            export_type="HISTORY_EXCEL", requested_by=self.user,
            params={"start": iso(T0), "end": iso(T_END), **extra},
        )

    def build(self, **extra):
        job = self.job(**extra)
        path = build_history_excel(job)
        return job, path, load_workbook(path)

    @staticmethod
    def rows(ws):
        return [[c for c in row] for row in ws.iter_rows(values_only=True)]


class HistoryExcelTests(ExcelBase):
    def test_sheets_headers_and_final_file(self):
        job, path, wb = self.build()
        self.assertEqual(path, output_path(job))
        self.assertTrue(path.is_file())
        self.assertFalse(list(path.parent.glob("*.part")))
        self.assertEqual(wb.sheetnames, ["Predictions", "Alerts", "WorkOrders", "Inventory"])
        for sheet, model in (("Predictions", Prediction), ("Alerts", Alert), ("WorkOrders", WorkOrder)):
            self.assertEqual(list(next(wb[sheet].iter_rows(values_only=True))), attnames(model), sheet)

    def test_predictions_range_start_inclusive_end_exclusive(self):
        _, _, wb = self.build()
        rows = self.rows(wb["Predictions"])
        id_col = attnames(Prediction).index("id")
        ids = {r[id_col] for r in rows[1:]}
        self.assertEqual(ids, {str(self.p_start.id), str(self.p_in1.id), str(self.p_in2.id)})

    def test_datetimes_are_naive_utc_and_decimals_numeric(self):
        _, _, wb = self.build()
        rows = self.rows(wb["Predictions"])
        header = rows[0]
        first = rows[1]  # diurutkan menaik menurut predicted_at -> p_start
        at = first[header.index("predicted_at")]
        self.assertIsNone(at.tzinfo)
        self.assertEqual(at, T0.replace(tzinfo=None))
        self.assertEqual(float(first[header.index("classification_confidence")]), 0.95)

    def test_unscoped_includes_null_bearing_alerts_and_all_wo_statuses_in_range(self):
        _, _, wb = self.build()
        alert_ids = {r[attnames(Alert).index("id")] for r in self.rows(wb["Alerts"])[1:]}
        self.assertEqual(alert_ids, {str(self.a_in.id), str(self.a_null.id)})
        wo_ids = {r[attnames(WorkOrder).index("id")] for r in self.rows(wb["WorkOrders"])[1:]}
        self.assertEqual(wo_ids, {str(self.wo_in.id)})

    def test_machine_scope_filters_every_sheet(self):
        _, _, wb = self.build(machine_ids=[str(self.m1.id)])
        pred_ids = {r[0 + attnames(Prediction).index("id")] for r in self.rows(wb["Predictions"])[1:]}
        self.assertEqual(pred_ids, {str(self.p_start.id), str(self.p_in1.id)})          # p_in2 milik mesin 2
        alert_ids = {r[attnames(Alert).index("id")] for r in self.rows(wb["Alerts"])[1:]}
        self.assertEqual(alert_ids, {str(self.a_in.id)})                                 # LOW_STOCK tanpa bearing tak ikut

    def test_bearing_scope(self):
        _, _, wb = self.build(bearing_ids=[str(self.b2.id)])
        pred_ids = {r[attnames(Prediction).index("id")] for r in self.rows(wb["Predictions"])[1:]}
        self.assertEqual(pred_ids, {str(self.p_in2.id)})
        self.assertEqual(len(self.rows(wb["Alerts"])), 1)       # hanya header

    def test_inventory_snapshot_and_movements_in_range(self):
        _, _, wb = self.build()
        rows = self.rows(wb["Inventory"])
        self.assertTrue(str(rows[0][0]).startswith("Spare parts"))
        self.assertEqual(trim(rows[1]), attnames(SparePart))
        self.assertEqual(rows[2][attnames(SparePart).index("sku")], "XL-1")      # snapshot tidak dibatasi rentang
        title_idx = next(i for i, r in enumerate(rows) if str(r[0] or "").startswith("Stock movements"))
        self.assertEqual(trim(rows[title_idx + 1]), attnames(StockMovement))
        movement_ids = {r[attnames(StockMovement).index("id")] for r in rows[title_idx + 2:]}
        self.assertEqual(movement_ids, {str(self.mv_in.id)})                       # mv_out di luar rentang

    def test_formula_text_is_stored_as_text_not_formula(self):
        _, _, wb = self.build()
        ws = wb["WorkOrders"]
        header = [c.value for c in ws[1]]
        notes = ws[2][header.index("notes")]
        self.assertEqual(notes.value, FORMULA)
        self.assertEqual(notes.data_type, "s")


class RunExportTaskTests(ExcelBase):
    def test_pending_to_completed(self):
        job = self.job()
        run_export(str(job.id))
        job.refresh_from_db()
        self.assertEqual(job.status, "COMPLETED")
        self.assertIsNotNone(job.completed_at)
        self.assertEqual(job.file_path, str(output_path(job)))
        self.assertTrue(output_path(job).is_file())

    def test_generator_failure_marks_failed_without_file_or_completed_at(self):
        job = self.job()
        with mock.patch("apps.exports.excel_report.build_history_excel", side_effect=RuntimeError("boom")):
            with self.assertLogs("apps.exports.tasks", level="ERROR"):
                run_export(str(job.id))
        job.refresh_from_db()
        self.assertEqual(job.status, "FAILED")
        self.assertIsNone(job.completed_at)
        self.assertIsNone(job.file_path)

    def test_non_pending_and_unknown_jobs_are_skipped(self):
        for status in ("RUNNING", "COMPLETED", "FAILED"):
            job = self.job()
            ExportJob.objects.filter(pk=job.pk).update(status=status)
            with mock.patch("apps.exports.excel_report.build_history_excel") as gen:
                run_export(str(job.id))
            gen.assert_not_called()
            job.refresh_from_db()
            self.assertEqual(job.status, status)
        run_export("00000000-0000-0000-0000-000000000000")  # tidak melempar
